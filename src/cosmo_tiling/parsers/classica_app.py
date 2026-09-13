"""Adapt the shared Classica column parser to app workbook rows."""
from collections import OrderedDict

from .classica_columns import (
    application_kind,
    extract,
    field,
    normalized,
    product_fields,
    structured_rows,
)
from .classica_rules import rules_for
from .common import OrderRow


def build_classica_document(path, template, *, apply_revisions=True):
    rules = rules_for(template)
    review_types = set(rules["review"]["row_types"])
    document = extract(path, apply_revisions=apply_revisions, rules=rules)
    rooms = OrderedDict()
    for raw in document["rows"]:
        rooms.setdefault(raw["room"], []).append(raw)
    rows = []
    for room, raw_rows in rooms.items():
        # Match rendered selections back to their original text and pattern.
        selections = {}
        pattern = ""
        for raw in raw_rows:
            if raw.get("excluded_by_change"):
                continue
            text = " ".join(raw["type_description"].split())
            if raw.get("layout", {}).get("is_heading") and application_kind(text, rules):
                pattern = ""
            if normalized(text).startswith("pattern:"):
                pattern = text.split(":", 1)[1].strip()
            if field(text, rules)[0] == "selection":
                selections.setdefault(product_fields(text, rules), []).append((raw, pattern))
        related = ""
        for rendered in structured_rows(raw_rows, rules):
            item_type, size, description = rendered["type"], rendered["size"], rendered["description"]
            canonical = application_kind(item_type, rules)
            pattern = rendered["pattern"]
            source = rendered["source"] or description
            if canonical:
                presentation = rules["applications"]["presentation"].get(canonical)
                if presentation:
                    display = item_type if presentation.get("preserve_source_type") else presentation["display"]
                    related = presentation["related"]
                else:
                    display, related = item_type, item_type
            else:
                display = item_type
                application_presentation = rules["applications"]["presentation"].get(
                    rendered["application"])
                if application_presentation:
                    related = application_presentation["related"]
            matches = selections.get((size, description), [])
            if matches:
                raw, pattern = matches[0]
                source = f"{raw['type_description']} | Qty: {raw['qty']} | Changed: {raw['selection_changed']}"
            elif item_type in {
                    rules["accessories"]["display"]["caulk"],
                    rules["accessories"]["drain_riser"]["type"],
                    rules["accessories"]["sealer"]["type"]}:
                source = f"Derived accessory rule for {room}: {description}"
            else:
                candidates = [r for r in raw_rows if description and description in r["type_description"]]
                if candidates:
                    source = candidates[0]["type_description"]
            rows.append(OrderRow(room=room, item_type=display, size_area=size,
                                 description=description,
                                 measured_qty=rendered["measured_qty"],
                                 order_qty=(rendered["quantity"] if rendered["quantity"] != "" else None),
                                 unit=rendered["unit"], pattern=pattern, source_text=source,
                                 waste_percent=rendered["waste_percent"],
                                 order_formula_override=rendered["formula"],
                                 related_type=related, comments=rendered["comments"]))
    for change in document["change_report"]:
        if change["status"] == "review":
            rows.append(OrderRow(room=change.get("room", "Review"),
                                 item_type=rules["review"]["required_type"],
                                 description=change["text"], comments=change["reason"],
                                 source_text=f"Change Order #{change['order']}: {change['text']}"))
    if not any(row.item_type not in review_types for row in rows):
        raise ValueError("No recognized tile-order selections were found in the PDF")
    return rows, document


def append_source_sheets(output_path, document, prefix=""):
    """Keep exact extracted values and revision decisions available for auditing."""
    from openpyxl import load_workbook
    workbook = load_workbook(output_path)
    source = workbook.create_sheet(prefix + "Source")
    source.append(["Room Code", "Type / Description", "Selection Changed", "Qty", "Page", "Excluded by change"])
    for raw in document["rows"]:
        values = [raw["room"], raw.get("original_description", raw["type_description"]),
                  raw["selection_changed"], raw["qty"], raw.get("layout", {}).get("page"),
                  raw.get("excluded_by_change", "")]
        source.append(values)
    source.freeze_panes = "A2"
    source.column_dimensions["B"].width = 100
    if document["change_report"]:
        report = workbook.create_sheet(prefix + "Change Review")
        report.append(["Change Order", "Status", "Room", "Application", "Instruction", "Reason"])
        for entry in document["change_report"]:
            report.append([entry["order"], entry["status"], entry.get("room", ""),
                           entry.get("application", ""), entry["text"], entry["reason"]])
        report.column_dimensions["E"].width = 100
    # PDF content is audit text, including any leading equals signs.
    for sheet in workbook.worksheets:
        if sheet.title in {prefix + "Source", prefix + "Change Review"}:
            for cells in sheet.iter_rows():
                for cell in cells:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
    workbook.save(output_path)
    workbook.close()
