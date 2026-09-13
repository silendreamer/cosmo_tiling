import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from cosmo_tiling.converter import load_template
from cosmo_tiling.parsers.classica_app import (
    REVIEW_TYPES, append_source_sheets, build_classica_document,
)
from cosmo_tiling.parsers.classica_columns import table_rows

ROOT = Path(__file__).resolve().parents[1]
PDF = ROOT / "pdf/classica/VendorOrder_PalosVerdeEstates7.pdf"
TEMPLATE = ROOT / "src/cosmo_tiling/config/templates/classica-template.json"


class ClassicaAppTests(unittest.TestCase):
    def test_app_preserves_shared_rules_and_dynamic_rooms(self):
        rows, document = build_classica_document(PDF, load_template(TEMPLATE))
        for room in {r["room"] for r in document["rows"]}:
            expected = table_rows([r for r in document["rows"] if r["room"] == room])
            actual = [r for r in rows if r.room == room and r.item_type != "Review required"]
            self.assertEqual(
                [(r.size_area, r.description, r.order_qty, r.unit) for r in actual],
                [(size, desc, qty if qty != "" else None, unit)
                 for _, size, desc, qty, unit in expected],
            )
        drains = 0
        for index, row in enumerate(rows):
            if row.item_type == "Shower Drain":
                drains += 1
                plug = rows[index + 1]
                self.assertEqual((plug.room, plug.item_type, plug.order_qty, plug.unit),
                                 (row.room, "Drain riser plug", 1, "EA"))
            elif row.item_type != "Drain riser plug":
                self.assertIsNone(row.order_qty)
        self.assertGreater(drains, 0)
        self.assertNotIn("projects", load_template(TEMPLATE))

    def test_source_sheets_preserve_literal_pdf_text(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "order.xlsx"
            Workbook().save(path)
            append_source_sheets(path, {"rows": [{"room": "NEW_ROOM",
                "type_description": "=literal", "selection_changed": "", "qty": "1 EA"}],
                "change_report": []}, prefix="Revised ")
            workbook = load_workbook(path)
            try:
                self.assertEqual(workbook["Revised Source"]["B2"].value, "=literal")
                self.assertEqual(workbook["Revised Source"]["B2"].data_type, "s")
            finally:
                workbook.close()
